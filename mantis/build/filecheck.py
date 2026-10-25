"""Check the map's links against a SharePoint file list export, and list the drawings the map doesn't link to.

    python3 build/filecheck.py <export.xlsx> data/mantis_v3c.json out/Mantis_file_check.xlsx

The export is SharePoint's "Export to Excel" of Shared Documents > Sunshine (all items, no folders view):
columns Name, Revision, Modified, Item Type, Path.
Prints use_address entries for links that can be fixed automatically (a file of the same name exists in ~Assets).
"""
import sys, json, re, collections, urllib.parse as up
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SITE = "https://vlineau.sharepoint.com/"
ASSETS = "sites/SCMantis/Shared Documents/Sunshine/Drawings/~Assets"
PURPLE = "6B1F7C"

rows = list(load_workbook(sys.argv[1], read_only=True).worksheets[0].iter_rows(values_only=True))
hdr = [str(h).strip().lower() for h in rows[0]]
col = {k: hdr.index(k) for k in ("name", "modified", "item type", "path")}
files, folders = {}, set()
for r in rows[1:]:
    name, path = r[col["name"]], r[col["path"]]
    if not name or not path:
        continue
    full = f"{path}/{name}"
    if r[col["item type"]] == "Folder":
        folders.add(full.lower())
    else:
        files[full.lower()] = dict(name=name, path=path, mod=r[col["modified"]])
by_name = collections.defaultdict(list)
for f in files.values():
    by_name[f["name"].lower()].append(f)
squash = lambda s: re.sub(r"[^a-z0-9.]", "", s.lower())
by_squashed = collections.defaultdict(list)
for f in files.values():
    by_squashed[squash(f["name"])].append(f)

d = json.load(open(sys.argv[2]))
OUT = sys.argv[3]
L = d["links"]


def spath(u):
    """Library path of a SharePoint link, or None if it isn't a plain file address."""
    if not u.startswith(SITE) or u.startswith(SITE + ":"):
        return None
    return up.unquote(up.unquote(u.split("?")[0]))[len(SITE):].replace("//", "/").rstrip("/")


def asset_url(f):
    return SITE + up.quote(f"{f['path']}/{f['name']}")


# ---------------- 1. every link
broken, fixes, not_checked = [], {}, collections.Counter()
linked = set()
for l in L:
    p = spath(l["u"])
    if p is None:
        not_checked["Google Maps directions" if "google." in l["u"] else "SharePoint sharing link (can't be checked from a file list)"] += 1
        continue
    if not p.lower().startswith("sites/scmantis/shared documents/sunshine/"):
        not_checked["File outside the Sunshine folder"] += 1
        continue
    if p.lower() in files:
        linked.add(p.lower())
        continue
    name = p.split("/")[-1]
    if name.lower() == ".pdf":
        broken.append([l.get("a") or "", l["t"], l.get("l") or "", p, "Link has no file name", ""])
        continue
    # same file name in ~Assets, or the same name with spaces/case squashed (e.g. "MW210G 65" -> MW210G65)
    cand = [f for f in by_name.get(name.lower(), []) if f["path"] == ASSETS] or \
           [f for f in by_squashed.get(squash(name), []) if f["path"] == ASSETS]
    if len(cand) == 1:
        f = cand[0]
        fixes[name.lower()] = asset_url(f)
        linked.add(f"{f['path']}/{f['name']}".lower())
        broken.append([l.get("a") or "", l["t"], l.get("l") or "", p, f"Fixed: now opens ~Assets/{f['name']}", ""])
    else:
        broken.append([l.get("a") or "", l["t"], l.get("l") or "", p, "No file with this name in SharePoint", ""])
# one row per distinct broken address
seen, uniq = set(), []
for b in broken:
    if b[3] not in seen:
        seen.add(b[3]); uniq.append(b + [sum(1 for x in broken if x[3] == b[3])])
broken = uniq

# ---------------- 2. drawings in ~Assets the map doesn't link to
ids = set(d["names"]) | {l["a"] for l in L if l.get("a")}
asset_files = collections.defaultdict(set)
for l in L:
    p = spath(l["u"])
    if l.get("a") and p:
        asset_files[l["a"]].add(p.split("/")[-1].rsplit(".", 1)[0].upper())
GENERAL = re.compile(r"^(CT_|BP_|SAP_)|network|architecture|distribution|flowchart|database|newdocument|geelong|electrologix|telephone", re.I)


def match_asset(stem):
    best = None
    for a in ids:
        A = a.upper().replace(" ", "")
        if len(A) >= 3 and (stem.startswith(A) or A.endswith(stem)) and (best is None or len(A) > len(best)):
            best = a
    return best


unlinked = []
for k, f in sorted(files.items()):
    if f["path"] != ASSETS or k in linked:
        continue
    stem = f["name"].rsplit(".", 1)[0].upper().replace(" ", "")
    a = None if GENERAL.search(f["name"]) else match_asset(stem)
    if GENERAL.search(f["name"]):
        kind, sug = "General document", "Not tied to one asset. Link it from the tool box if it's useful, otherwise leave it."
    elif a is None:
        kind, sug = "Asset not on this map", "Probably an asset outside the Sunshine map, or retired. Leave, or move/delete if it's a stray."
    else:
        dup = [s for s in asset_files[a] if s != stem and (s.endswith(stem) or stem.endswith(s))]
        if dup:
            kind, sug = "Possible duplicate", f"{a} already links to {dup[0]}.pdf. If they're the same drawing, delete one."
        elif re.search(r"G(65|80)", stem):
            kind, sug = "Speed indicator drawing", f"The map's speed box at {a} opens the signal drawing. Could open this instead."
        else:
            kind, sug = "Drawing for an asset on the map", f"{a} is on the map but nothing links to this. Add it to {a}'s panel?"
    unlinked.append([f["name"], kind, a or "", sug, f["mod"], ""])
ORDER = ["Speed indicator drawing", "Drawing for an asset on the map", "Possible duplicate", "Asset not on this map", "General document"]
unlinked.sort(key=lambda r: (ORDER.index(r[1]), r[0].upper()))

# ---------------- 3. asset list
alist = []
for a in sorted(ids, key=str.upper):
    ls = [l for l in L if l.get("a") == a]
    types = collections.Counter(l["t"] for l in ls)
    docs, missing = set(), set()
    for l in ls:
        p = spath(l["u"])
        if p is None:
            continue
        nm = p.split("/")[-1]
        docs.add(nm)
        if p.lower() not in files and nm.lower() not in fixes:
            missing.add(nm)
    extra = [r[0] for r in unlinked if r[2] == a]
    alist.append([a, types.most_common(1)[0][0] if types else "", "Yes" if a in d["names"] else "Symbol only",
                  ", ".join(sorted(docs)), ", ".join(sorted(missing)) or "", ", ".join(extra)])

# ---------------- workbook
HDR, FILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=PURPLE)
WRAP, thin = Alignment(wrap_text=True, vertical="top"), Side(style="thin", color="D0D7DE")


def table(ws, headers, data, widths, answer=None):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR, FILL, Alignment(vertical="center", wrap_text=True)
    for r in data:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment, c.border = WRAP, Border(bottom=thin)
            if hasattr(c.value, "year"):
                c.number_format = "dd mmm yyyy"
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    if answer and data:
        dv = DataValidation(type="list", formula1=answer, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(len(headers) if headers[-1] == 'Your answer' else len(headers) - 1)}2:"
               f"{get_column_letter(len(headers) if headers[-1] == 'Your answer' else len(headers) - 1)}{len(data) + 1}")


wb = Workbook()
ws = wb.active
ws.title = "Read me"
n_ok = sum(1 for l in L if (spath(l["u"]) or "").lower() in files)
kinds = collections.Counter(r[1] for r in unlinked)
info = [["Mantis file check"], [""],
        ["Files in the SharePoint export", len(files)],
        ["Map links checked against it", n_ok + sum(r[-1] for r in broken)],
        ["  open a real file", n_ok],
        ["  fixed automatically (file is in ~Assets)", sum(r[-1] for r in broken if r[4].startswith("Fixed"))],
        ["  still broken", sum(r[-1] for r in broken if not r[4].startswith("Fixed"))]]
info += [[f"Not checked: {k}", v] for k, v in not_checked.items()]
info += [[""], ["Drawings in ~Assets the map doesn't link to", len(unlinked)]] + [[f"  {k}", kinds[k]] for k in ORDER if kinds[k]]
info += [[""], ["Sheets"],
         ["Broken links", "Links that don't open a real file, one row per address. 'Fixed' rows need nothing from you."],
         ["Unlinked drawings", "Files in ~Assets that nothing on the map opens, sorted by what's most likely worth linking."],
         ["Asset list", "Every asset on the map, the drawings it opens, any that are missing, and unlinked drawings that look like its."]]
for r in info:
    ws.append(r)
ws["A1"].font = Font(bold=True, size=15, color=PURPLE)
ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 52, 100
for row in ws.iter_rows(min_row=3):
    for c in row:
        c.alignment = WRAP

table(wb.create_sheet("Broken links"), ["Asset", "Type", "Text on map", "Link points to", "Status", "Your answer", "Links"],
      broken, [14, 22, 18, 70, 40, 30, 8], '"Leave it,Remove the link,File name in note"')
table(wb.create_sheet("Unlinked drawings"), ["File", "What it looks like", "Asset", "Suggestion", "Modified", "Your answer"],
      unlinked, [30, 26, 14, 64, 14, 30], '"Link it,Leave it,Delete from SharePoint,Not sure"')
table(wb.create_sheet("Asset list"), ["Asset", "Type", "Name on map", "Drawings it opens", "Missing in SharePoint", "Unlinked drawings that look like its"],
      alist, [14, 22, 14, 60, 30, 40])
wb.save(OUT)

print("links ok", n_ok, "broken rows", len(broken), "not checked", dict(not_checked))
print("unlinked", len(unlinked), dict(kinds))
print("use_address:")
print(json.dumps(fixes, indent=2))
