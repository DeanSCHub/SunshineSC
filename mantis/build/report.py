"""Mantis link check workbook."""
import sys, json, re, collections, urllib.parse as up
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

d = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
L = d["links"]
names = d["names"]
by_i = {l["i"]: l for l in L}
PURPLE = "6B1F7C"
HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor=PURPLE)
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D0D7DE")


def fname(u):
    try:
        return up.unquote(up.unquote(u.split("?")[0])).rstrip("/").split("/")[-1]
    except Exception:
        return u


def table(ws, headers, rows, widths, freeze="A2"):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR, FILL, Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            c.border = Border(bottom=thin)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


# ---------------- checks needing a person
check = []
issues = d["issues"]
for x in issues:
    if x["kind"] == "Two links stacked on one symbol":
        lose = by_i.get(x["i"], {})
        win = by_i.get(lose.get("x"), {})
        where = win.get("a") or lose.get("a") or ""
        check.append(["Two links on one symbol", where,
                      f"{win.get('f', '')} and {lose.get('f', '')} sit on exactly the same spot in the PDF.",
                      f"Tapping it opens {where}. {fname(lose.get('u', ''))} is listed under 'Also linked here' in {where}'s panel.",
                      "Confirm which drawing belongs there. If both are right, nothing to do. If one is a copy-paste slip, delete it from the asset list."])
for x in issues:
    if x["kind"] == "Name on map with no link" and x["asset"] in ("SUN946", "SUN571"):
        check.append(["Signal with no drawing link", x["asset"], f"{x['asset']} is written on the map but nothing links to its drawing.",
                      "Tapping the number does nothing.", f"Confirm the file name (probably {x['asset']}G.pdf in ~Assets) so it can be added."])
for x in issues:
    if x["kind"] == "Name on map with no link" and x["asset"] == "MW325":
        check.append(["Tag with no link", "MW325 CHCPB", "The MW325 CHCPB tag has no link.", "Tapping it does nothing.",
                      "Add a link if there's a drawing for it, otherwise ignore."])
for x in issues:
    if x["kind"] == "Link opens a different asset":
        check.append(["Label opens a different drawing", x["text"], f"The link on the label {x['text']} opens {x['opens']}.",
                      f"Tapping {x['text']} opens {x['opens']}.", f"Probably should open {x['text']}.pdf. Confirm and correct."])
for x in issues:
    if x["kind"] == "Link has no file name":
        check.append(["Link with no file name", x["text"], f"The link on {x['text']} points at the ~Assets folder with no file name.",
                      "Shows 'Drawing not linked' instead of opening a broken page.", "Tell me the right file name."])
for x in issues:
    if x["kind"] == "Same drawing, different addresses":
        urls = x["note"].split("; ")
        check.append(["Same drawing in two places", x["asset"], "The same file is linked from two different locations:\n" + "\n".join(urls),
                      "The panel shows it once, preferring the ~Assets copy on the SCMantis site.",
                      "Check the two copies are identical and delete the stray one, so there's only one to keep up to date."])

# ---------------- fixed automatically
fixed = []
for f in d.get("fixes", []):
    if f["kind"].startswith("Fixed: removed double slash"):
        fixed.append(["Double slash in address", f["text"], f["before"], f["after"]])
for f in d.get("fixes", []):
    if f["kind"] == "Merged IDs":
        fixed.append(["One signal under two IDs", f["text"], "Signal drawing filed as e.g. DPWA201G.pdf, TPWS as A201TPWS.pdf",
                      "Treated as one asset under the ID written on the map (A201 etc.)"])
fixed.append(["Highlight covered the whole area", "All assets", "Selecting a signal highlighted a box around every link it had (number, TPWS tag, arrows).",
              "Highlight now sits only on the written name, e.g. just 'WVS933'."])
fixed.append(["Signal numbers weren't tappable", "About 170 signals", "In the PDF many signals only had a link on the triangle, not the number.",
              "Every written asset name is now a tap target as well as its symbol."])
fixed.append(["Duplicate drawings in the panel", "e.g. WVS933", "WVS933TPWS was listed twice (.PDF and .pdf).", "Each drawing now shows once."])
fixed.append(["Drawing bands on the map", "Whole map", "SAP, bonding plan and control table stripes", "Removed from the map. Each asset's panel lists 'Drawings for this area'."])

# ---------------- all links
stacked_losers = {l["i"] for l in L if "x" in l}
rows = []
for l in sorted(L, key=lambda l: (l["a"] or "~", l["t"], l["f"])):
    if l["t"] == "Directions":
        status = "OK (directions pin)"
    elif l.get("m"):
        status = "No file name"
    elif l["i"] in stacked_losers:
        status = f"Stacked under {by_i[l['x']].get('a') or by_i[l['x']]['f']}: shown as 'Also linked here'"
    elif re.search(r"/Drawings//", up.unquote(l["u"])):
        status = "Fixed"
    elif l["t"] in ("Signalling arrangement plan", "Bonding plan", "Control tables"):
        status = "OK (offered from assets in this area)"
    else:
        status = "OK"
    tap = "Name and symbol" if l["a"] and l["a"] in names else ("Symbol / label" if l["t"] not in ("Signalling arrangement plan", "Bonding plan", "Control tables", "Directions") else "From panel")
    rows.append([l["a"], l["t"], l.get("l", ""), fname(l["u"]), l["u"], tap, status])

wb = Workbook()
ws = wb.active
ws.title = "Read me"
total = len(L)
ok = sum(1 for r in rows if r[6].startswith("OK"))
info = [
    ["Mantis link check"],
    [""],
    ["Links in the PDF", total],
    ["Working as they should", ok],
    ["Fixed automatically", len(fixed)],
    ["Need a person to check", len(check)],
    [""],
    ["How this was checked"],
    ["", "Every link was matched against the words and symbols underneath it in the original PDF, then all 1,688 tap points (every asset name plus every link) were run through the map's tap logic at two zoom levels to confirm each one opens the right asset. All pass."],
    ["", "What this can't check: whether each SharePoint file actually exists. For that, send an export of the Drawings library file list (see 'What I need' below)."],
    [""],
    ["Sheets"],
    ["Needs checking", "Items only you can answer. There's a column for your answer."],
    ["Fixed automatically", "What was corrected in the new Mantis file. No action needed."],
    ["All links", "Every link, what it opens, how you reach it on the map, and its status."],
    [""],
    ["What I need"],
    ["1", "A file list of the Drawings library: in SharePoint open Shared Documents > Sunshine > Drawings, switch to a view that shows all items without folders, then Export > Export to CSV. With that I can confirm every link opens a real file and fix any that don't."],
    ["2", "Your answers in the 'Needs checking' sheet."],
]
for r in info:
    ws.append(r)
ws["A1"].font = Font(bold=True, size=15, color=PURPLE)
for c in ("A8", "A12", "A17"):
    ws[c].font = Font(bold=True)
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 110
for row in ws.iter_rows(min_row=3):
    for c in row:
        c.alignment = WRAP

ws2 = wb.create_sheet("Needs checking")
table(ws2, ["Issue", "Where", "What's wrong", "What Mantis does now", "Suggested fix", "Your answer"], check, [26, 16, 48, 40, 44, 34])
dv = DataValidation(type="list", formula1='"Fixed in SharePoint,Correct as is,Change it - see note,Not sure"', allow_blank=True)
ws2.add_data_validation(dv)
dv.add(f"F2:F{len(check) + 1}")
for r in range(2, len(check) + 2):
    ws2.row_dimensions[r].height = 75

ws3 = wb.create_sheet("Fixed automatically")
table(ws3, ["What", "Where", "Before", "After"], fixed, [30, 30, 60, 60])

ws4 = wb.create_sheet("All links")
table(ws4, ["Asset", "Type", "Text on map", "Opens", "Address", "Tap target", "Status"], rows, [14, 26, 22, 30, 70, 18, 38])
wb.save(OUT)
print("checks", len(check), "fixed", len(fixed), "rows", len(rows), "ok", ok)
